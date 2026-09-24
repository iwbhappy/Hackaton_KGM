"""Two complete acceptance runs, each on a newly created SQLite database."""
from io import BytesIO
from time import monotonic
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from app.config import ROOT
from tests.test_scanner_api import route_lab


@pytest.mark.parametrize("round_number", [1, 2])
def test_complete_demo_from_clean_database(round_number, client, monkeypatch, lab_server, lab_material):
    route_lab(monkeypatch, lab_server, lab_material)
    notifier = Mock()
    monkeypatch.setattr("app.notifiers.dispatcher.get_notifiers", lambda: {"telegram": notifier})
    assert client.get("/api/results").json() == []
    csv = (ROOT / "lab/targets.csv").read_bytes()
    assert client.post("/api/endpoints/import", files={"file": ("targets.csv", csv)}).status_code == 200
    demo = client.get("/api/demo-targets").json()
    start = monotonic()
    response = client.post("/api/scans", json=demo)
    assert response.status_code == 202
    assert monotonic() - start < 15
    rows = {row["host"]: row for row in client.get("/api/results").json()}
    assert len(rows) == 15
    assert rows["vpn.lab.local"]["risk_score"] == 91
    assert rows["vpn.lab.local"]["days_left"] == 5
    for host, code in [("nochain", "CHAIN_UNTRUSTED"), ("mismatch", "HOSTNAME_MISMATCH"),
                       ("weak", "WEAK_KEY"), ("weak", "WEAK_SIGNATURE"), ("dead", "UNREACHABLE")]:
        assert code in {i["code"] for i in rows[host + ".lab.local"]["issues"]}
    no_owner = rows["nochain.lab.local"]
    updated = client.patch(f"/api/endpoints/{no_owner['endpoint_id']}", json={"owner": "ИТ"}).json()
    assert updated["latest"]["risk_score"] == no_owner["risk_score"] - 10
    assert updated["latest"]["scanned_at"] == no_owner["scanned_at"]
    assert notifier.send.call_count == 1
    assert client.post("/api/scans", json={"rescan_all": True}).status_code == 202
    assert notifier.send.call_count == 1
    assert client.post("/api/notifications/send-now").status_code == 200
    assert notifier.send.call_count == 2
    book = load_workbook(BytesIO(client.get("/api/export/xlsx").content))
    assert len(book.sheetnames) == 3 and book.worksheets[0].max_row == 16
    assert client.get("/api/export/csv").content.startswith(b"\xef\xbb\xbf")
    assert "Рекомендации" in client.get("/report").text
    assert client.put("/api/settings", json={"threshold_info": 40}).status_code == 200
    info_id = rows["info.lab.local"]["endpoint_id"]
    assert client.get(f"/api/endpoints/{info_id}").json()["latest"]["status"] == "OK"
    assert client.put("/api/settings", json={"threshold_info": 60, "threshold_warning": 50}).status_code == 200
    assert client.get(f"/api/endpoints/{info_id}").json()["latest"]["status"] == "WARNING"
    events = {entry["action"] for entry in client.get("/api/audit").json()}
    assert {"app_started", "scan_started", "scan_finished", "endpoint_updated", "export_downloaded", "notification_sent", "settings_changed"} <= events
    assert len(client.get(f"/api/endpoints/{info_id}").json()["history"]) == 2


def test_fifty_tls_targets_under_fifteen_seconds(client, monkeypatch, lab_server, lab_material):
    import socket
    route_lab(monkeypatch, lab_server, lab_material)
    original = socket.getaddrinfo

    def resolve(host, port, *args, **kwargs):
        if host == "127.0.0.1" and 10000 <= port < 10050:
            port = lab_server.server_address[1]
        return original(host, port, *args, **kwargs)

    monkeypatch.setattr(socket, "getaddrinfo", resolve)
    targets = "\n".join(f"valid.lab.local:{port}" for port in range(10000, 10050))
    start = monotonic()
    assert client.post("/api/scans", json={"text": targets}).status_code == 202
    elapsed = monotonic() - start
    rows = client.get("/api/results").json()
    assert len(rows) == 50 and all(row["chain_status"] == "valid" for row in rows)
    assert elapsed < 15
    print(f"50 реальных TLS-целей: {elapsed:.2f} с")
