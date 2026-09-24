import csv
from io import BytesIO, StringIO
from zipfile import ZipFile

from openpyxl import load_workbook
from tests.test_endpoints import add_result


def test_all_exports_cyrillic_colors_and_safe_cells(client):
    endpoint_id=add_result(owner="Иванов И.", days=5)
    client.patch(f"/api/endpoints/{endpoint_id}", json={"service_name":"=2+2<script>alert(1)</script>"})
    response=client.get("/api/export/csv")
    assert response.content.startswith(b"\xef\xbb\xbf")
    rows=list(csv.reader(StringIO(response.content.decode("utf-8-sig")),delimiter=";"))
    assert rows[1][0].startswith("'=2+2")
    assert "Иванов И." in rows[1]
    response=client.get("/api/export/xlsx")
    assert ZipFile(BytesIO(response.content)).testzip() is None
    book=load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Сертификаты","Сводка","Проблемы"]
    sheet=book["Сертификаты"]
    assert sheet["A2"].data_type == "s"
    assert sheet["D2"].value == "Иванов И."
    assert sheet["J2"].fill.fgColor.rgb == "00FFF0E3"
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref
    report=client.get("/report").text
    assert "&lt;script&gt;" in report and "<script>alert(1)</script>" not in report
    assert "@media print" in report and "Рекомендации" in report
