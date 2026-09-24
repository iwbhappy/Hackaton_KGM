"""Three-sheet workbook with localized cells and status colors."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.exporters.common import COLUMNS, STATUS_COLORS, STATUS_LABELS, display_value
from app.repository import summarize


def append_text_row(sheet, values: list) -> None:
    """Write strings as literal text, including values beginning with '='."""
    sheet.append([ILLEGAL_CHARACTERS_RE.sub("�", item) if isinstance(item, str) else item for item in values])
    for cell, value in zip(sheet[sheet.max_row], values):
        if isinstance(value, str):
            cell.data_type = "s"


def style_sheet(sheet) -> None:
    """Apply readable widths, wrapping, fixed header and an autofilter."""
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="19354D")
        cell.font = Font(color="FFFFFF", bold=True)
    for column in sheet.columns:
        width = min(60, max(13, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_xlsx(rows: list[dict]) -> bytes:
    """Build certificate, summary and one-row-per-issue worksheets."""
    book = Workbook()
    certs = book.active
    certs.title = "Сертификаты"
    append_text_row(certs, [title for _, title in COLUMNS])
    status_col = [key for key, _ in COLUMNS].index("status") + 1
    for row in rows:
        append_text_row(certs, [display_value(row, key) for key, _ in COLUMNS])
        certs.cell(certs.max_row, status_col).fill = PatternFill("solid", fgColor=STATUS_COLORS[row["status"]])
    summary = book.create_sheet("Сводка")
    append_text_row(summary, ["Показатель", "Значение"])
    kpi = summarize(rows)
    for label, value in [("Сервисы", kpi["services"]), ("Уникальные сертификаты", kpi["certificates"]),
                          ("Здоровье инфраструктуры", kpi["health"]), *[(STATUS_LABELS[s], n) for s, n in kpi["statuses"].items()],
                          ("Проблемы цепочки", kpi["issues"]["CHAIN"]), ("Несоответствие имени", kpi["issues"]["HOSTNAME_MISMATCH"]),
                          ("Без владельца", kpi["issues"]["NO_OWNER"])]:
        append_text_row(summary, [label, value])
    problems = book.create_sheet("Проблемы")
    append_text_row(problems, ["Сервис", "Владелец", "Код", "Проблема", "Подробности", "Рекомендация", "Баллы"])
    for row in rows:
        for issue in row["issues"]:
            append_text_row(problems, [row["service_name"], row["owner"], issue["code"], issue["title"],
                                       issue["detail"], issue["recommendation"], issue["points"]])
    for sheet in book:
        style_sheet(sheet)
    output = BytesIO()
    book.save(output)
    return output.getvalue()
